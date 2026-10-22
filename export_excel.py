"""Write the results workbook."""
import logging
from pathlib import Path

import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def _format_sheet(ws, df: pd.DataFrame) -> None:
    for idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        series = df[col]
        if pd.api.types.is_datetime64_any_dtype(series):
            fmt = "mmm-yyyy" if col == "month" else "dd-mm-yyyy"
        elif pd.api.types.is_float_dtype(series):
            fmt = "#,##0.00"
        else:
            fmt = None
        sample = [len(str(col))] + [len(str(v)) for v in series.head(200)]
        ws.column_dimensions[get_column_letter(idx)].width = min(max(sample) + 3, 28)
        if fmt:
            for row in range(2, len(df) + 2):
                ws.cell(row=row, column=idx).number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write(writer, df: pd.DataFrame, name: str) -> None:
    df.to_excel(writer, sheet_name=name[:31], index=False)
    _format_sheet(writer.sheets[name[:31]], df)


def export_workbook(path: Path, monthly: pd.DataFrame, daily: pd.DataFrame,
                    weather_monthly: pd.DataFrame, notes: list[str]) -> None:
    if not weather_monthly.empty:
        monthly = monthly.merge(weather_monthly, on="month", how="left")

    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="dd-mm-yyyy") as writer:
        _write(writer, monthly, "Monthly Summary")
        for seg, part in daily.groupby("segment"):
            _write(writer, part.reset_index(drop=True), f"{seg} Daily")
        if not weather_monthly.empty:
            _write(writer, weather_monthly, "Weather Monthly")

        pivot = monthly.pivot(index="month", columns="segment", values="mcp_avg").reset_index()
        _write(writer, pivot, "MCP by Month")
        ws = writer.sheets["MCP by Month"]
        chart = LineChart()
        chart.title = "Average MCP by month (Rs/MWh)"
        chart.y_axis.title = "Rs/MWh"
        chart.x_axis.number_format = "mmm-yy"
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        n_rows, n_cols = len(pivot), pivot.shape[1]
        chart.add_data(Reference(ws, min_col=2, max_col=n_cols, min_row=1, max_row=n_rows + 1), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n_rows + 1))
        chart.height, chart.width = 9, 22
        ws.add_chart(chart, f"{get_column_letter(n_cols + 2)}2")

        notes_df = pd.DataFrame({"Notes": notes})
        notes_df.to_excel(writer, sheet_name="Notes", index=False)
        writer.sheets["Notes"].column_dimensions["A"].width = 120
    log.info("Wrote %s", path)
